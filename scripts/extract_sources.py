import json,re
from pathlib import Path
import openpyxl

ROOT=Path(__file__).resolve().parent.parent
WORKBOOK=ROOT/'营运拓展营拓技能清单-TO庚涛.xlsx'
wb=openpyxl.load_workbook(WORKBOOK,data_only=True)
position_skills=json.loads((ROOT/'src'/'position-skills.json').read_text())

def merged_value(ws,cell):
 if cell.value is not None:return cell.value
 for rng in ws.merged_cells.ranges:
  if cell.coordinate in rng:return ws.cell(rng.min_row,rng.min_col).value
 return None

def code_from(value):
 match=re.match(r'^([YSTM]\d+(?:-\d+)?)',str(value or '').strip())
 return match.group(1) if match else ''

skill_names={
 'Y1':'执行巡店','Y2':'发现门店问题','Y3':'改善门店问题',
 'Y4-1':'改善区域问题（复核与推动）','Y4-2':'改善区域问题（判断与落地）','Y4-3':'经营诊断分析','Y5':'制定专项',
 'S1':'加盟商协同','S3':'加盟商在店判断','S4-1':'加盟商综合判断（提出建议）','S4-2':'加盟商综合判断（方案判断）','S4-3':'加盟商综合判断（发展规划）','S5':'加盟商发展规划',
 'T1':'摸排与单点盈利测算','T2':'点位带看','T3':'单点判断','T4':'最佳建店方案','T5':'多点位商圈规划','T6':'区县规划',
 'M1':'带教复核','M2':'梯队建设'
}
group_names={'店':'门店营运','人':'加盟商发展','点':'点位拓展','团队':'团队管理'}

# 认证方式只补充课程和验证标准，不决定岗位需要哪些技能。
cert=wb['认证方式'];cert_index={}
for row in range(2,cert.max_row+1):
 axis=str(merged_value(cert,cert.cell(row,2)) or '')
 code=code_from(merged_value(cert,cert.cell(row,3)))
 description=str(cert.cell(row,4).value or '').strip()
 if not description:continue
 if '商' in axis and code.startswith('Y'):code='S'+code[1:]
 if code:
  cert_index.setdefault(code,[]).append({
   'description':description,
   'course':str(merged_value(cert,cert.cell(row,5)) or '').strip(),
   'verification':str(merged_value(cert,cert.cell(row,6)) or '').strip(),
   'source':f'认证方式!C{row}:F{row}'
  })

# 三张明细表只补充评估方式；评估方式按岗位序列保存，不能跨序列直接拼接。
assessment_index={}
for sheet_name,track in [('纯营技能清单','营运'),('营拓技能清单','营拓'),('纯拓技能清单','拓展')]:
 ws=wb[sheet_name]
 for row in range(3,ws.max_row+1):
  code=code_from(merged_value(ws,ws.cell(row,3)))
  description=str(ws.cell(row,4).value or '').strip()
  assessment=str(ws.cell(row,6).value or '').strip()
  if code and assessment and assessment!='/':
   assessment_index.setdefault(code,{}).setdefault(track,[]).append((description,assessment))

def assessment_text(items):
 unique=list(dict.fromkeys(items))
 methods=list(dict.fromkeys(method for _,method in unique))
 # 同一技能编码在明细表中可能跨多行合并；技能名称对应首行定义，
 # 后续不同方法属于下一子项，不能拼成同一条认证方式。
 return methods[0]

unique={}
for items in position_skills.values():
 for item in items:
  code=item['code'];axis=item['axis'];key=code
  if key not in unique:
   cert_matches=cert_index.get(code,[])
   base=re.sub(r'-\d+$','',code)
   if not cert_matches:cert_matches=cert_index.get(base,[])
   course='\n'.join(dict.fromkeys(x['course'] for x in cert_matches if x['course'] and x['course']!='/')) or '课程待补充'
   verification='\n'.join(dict.fromkeys(x['verification'] for x in cert_matches if x['verification'])) or '结果验证标准待补充'
   assessment_variants={track:assessment_text(items) for track,items in assessment_index.get(base,{}).items()}
   assessment=next(iter(dict.fromkeys(assessment_variants.values())),'认证方式待补充')
   unique[key]={
    'id':'skill-'+code.lower(),
    'code':code,
    'name':skill_names.get(code,code),
    'group':group_names[axis],
    'definition':item['description'],
    'course':course,
    'assessment':assessment,
    'assessments':assessment_variants,
    'verification':verification,
    'source':item['source']+(('；认证方式｜'+cert_matches[0]['source']) if cert_matches else '')
   }

requirements=[]
for position_id,items in position_skills.items():
 for item in items:
  requirements.append({
   'positionId':position_id,
   'skillId':'skill-'+item['code'].lower(),
   'level':1,
   'kind':'必备',
   'raw':item['description'],
   'source':item['source'],
   'confirmed':True,
   'comparable':False
  })

output={
 'skills':list(unique.values()),
 'requirements':requirements,
 'rawMatrixRows':[],
 'notes':[
  '岗位与技能代码的对应关系只取自“技能清单”Sheet。',
  '“认证方式”Sheet 只用于补充课程和结果验证标准。',
  '三张分路径技能明细表只用于补充评估方式，不参与岗位技能映射。',
  '软性素质按岗位统一概述，不作为逐项认证技能。'
 ]
}
(ROOT/'src'/'source-data.json').write_text(json.dumps(output,ensure_ascii=False,indent=2))
print(len(unique),'skills',len(requirements),'requirements from 技能清单')

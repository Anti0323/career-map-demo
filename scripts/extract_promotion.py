import json
from pathlib import Path
import openpyxl
root=Path(__file__).resolve().parent.parent
s=openpyxl.load_workbook(root/'拓展晋升标准-TO庚涛.xlsx',data_only=True).active
rules=[]
for col,applies in [(2,[('拓展',5,6)]),(3,[('拓展',6,7),('营拓',7,8)]),(4,[('拓展',7,8),('营拓',8,9)]),(5,[('拓展',8,9),('营拓',9,10)])]:
 letter=openpyxl.utils.get_column_letter(col)
 for track,frm,to in applies:
  rules.append({'id':f'{track}-{frm}-{to}','track':track,'fromLevel':frm,'toLevel':to,'title':str(s.cell(14,col).value).split('\n')[0],'nomination':str(s.cell(11,col).value or ''),'reviewPeriod':str(s.cell(12,col).value or ''),'annualLimit':str(s.cell(13,col).value or ''),'passStandard':str(s.cell(14,col).value or ''),'source':f'拓展晋升标准-TO庚涛.xlsx｜Sheet1!{letter}11:{letter}14','confirmed':False if track=='营拓' else True})
(root/'src'/'promotion-data.json').write_text(json.dumps(rules,ensure_ascii=False,indent=2))
print(len(rules),'transition rules')
